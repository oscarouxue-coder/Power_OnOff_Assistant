import os
import re
import openpyxl


def get_latest_file_method(path):
    # 获取目录下所有文件
    files = os.listdir(path)
    # 按照修改时间排序
    files.sort(key=lambda x: os.path.getmtime(os.path.join(path, x)))
    # 获取最新修改时间的文件
    latest_file = files[-1]
    # 返回最新修改时间和最新修改时间文件的名称
    return os.path.getmtime(os.path.join(path, latest_file)), latest_file


if __name__ == "__main__":
    path = r'E:\Gen7_Projects_04 _v1.2.3\test_results'
    latest_time, latest_file_name = get_latest_file_method(path)
    # print('最新修改时间：', latest_time)
    print('最新修改时间文件的名称：', latest_file_name)

    mac_re = re.compile('address (.*) under')

    file_path = path + '\\' + latest_file_name

    workbook = openpyxl.load_workbook('../PCIE_result_statistics.xlsx')
    sheets = workbook.sheetnames
    ws = workbook[sheets[0]]

    max_row = ws.max_row
    max_col = ws.max_column
    print(max_row, max_col)
    # 获取MAC列的值
    mac_list = []
    for i in range(max_row):
        mac_list.append(ws.cell(i + 1, 2).value)
    print(mac_list)
    # row = 2
    # col = 1
    position_number = 1
    with open(file_path, 'rb') as f:
        line = f.readline().decode()
        status_code = line.split()[-1].split(':')[-1]
        print(status_code)
        mac = mac_re.search(line).group(1)
        print(mac_re.search(line).group(1))
        ws.cell(max_row + 1, column=1, value=max_row)

        # 判断MAC的是否已经在excel表格里面
        # count = 1
        # mac_flag = 0
        # for i in mac_list:
        #     if mac == i:
        #         mac_flag = 1
        #         break
        #     else:
        #         count += 1
        # print('count', count)
        # if mac_flag:
        #     ws.cell(count+1, 3, value=ws.cell(count, column=3).value + 1)

        if ws.cell(max_row + 1, column=3).value == None:
            ws.cell(max_row + 1, column=3, value=1)
        else:
            ws.cell(max_row + 1, column=3, value=ws.cell(max_row + 1, column=3).value + 1)
        workbook.save('../PCIE_result_statistics.xlsx')
