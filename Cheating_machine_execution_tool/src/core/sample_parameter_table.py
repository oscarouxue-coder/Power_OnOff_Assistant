from __future__ import annotations

import csv
import secrets
from pathlib import Path
from typing import Iterable

from src.models import (
    PARAMETER_COLUMNS,
    STANDARD_PARAMETER_DEFINITIONS,
    VLAN_COLUMN,
    ParameterGenerationTemplate,
    SampleParameterRecord,
)


XPON_SN_TOTAL_LENGTH = 12
XPON_SN_HEX_SUFFIX_LENGTH = 8
XPON_SN_PREFIX_LENGTH = XPON_SN_TOTAL_LENGTH - XPON_SN_HEX_SUFFIX_LENGTH
_UPPER_HEX_ALPHABET = "0123456789ABCDEF"
SN_MAX_LENGTH = 16
_DIGIT_ALPHABET = "0123456789"


def load_parameter_records(path: Path) -> list[SampleParameterRecord]:
    rows = _read_rows(path)
    if not rows:
        raise ValueError("参数表为空。")
    header = [str(value).strip() for value in rows[0]]
    if not _indices_for(header, VLAN_COLUMN):
        raise ValueError("参数表缺少必填列 VLAN。")
    records: list[SampleParameterRecord] = []
    for row_number, raw_row in enumerate(rows[1:], start=2):
        row = ["" if value is None else str(value).strip() for value in raw_row]
        if not any(row):
            continue
        vlan = _cell_for(header, row, VLAN_COLUMN)
        values = {
            definition.column: _cell_for(header, row, definition.column, definition.legacy_columns)
            for definition in STANDARD_PARAMETER_DEFINITIONS
        }
        records.append(SampleParameterRecord(vlan=vlan, values=values, source_row=row_number))
    return records


def save_parameter_records_xlsx(path: Path, records: Iterable[SampleParameterRecord]) -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("导出 .xlsx 需要安装 openpyxl。") from exc
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "样机参数"
    headers = (VLAN_COLUMN, *PARAMETER_COLUMNS)
    sheet.append(headers)
    for record in records:
        sheet.append([record.vlan, *(record.values.get(column, "") for column in PARAMETER_COLUMNS)])
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 30)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def generate_parameter_records(vlans: Iterable[str], template: ParameterGenerationTemplate) -> list[SampleParameterRecord]:
    suffix_length = int(template.suffix_length)
    if not 1 <= suffix_length <= SN_MAX_LENGTH:
        raise ValueError(f"SN 随机后缀长度必须在 1 到 {SN_MAX_LENGTH} 之间。")
    normalized_vlans = [str(vlan).strip() for vlan in vlans if str(vlan).strip()]
    if not normalized_vlans:
        raise ValueError("没有可生成的 VLAN。")
    if len({item.lower() for item in normalized_vlans}) != len(normalized_vlans):
        raise ValueError("已勾选通道存在重复 VLAN，无法生成参数表。")
    seen_mac: set[str] = set()
    seen_values: dict[str, set[str]] = {"SN": set(), "GPON SN": set()}
    # 同批 MAC 前 3 字节由用户指定，后 3 字节为每台唯一的随机十六进制值。
    mac_prefix = _normalize_mac_prefix(template.mac_prefix)
    # 2.4G / 5G SSID 前缀必须非空(用于拼接 MAC 末 4 位作为共同后缀);
    # 若用户没填,与旧版"SSID 前缀不能为空"的约束保持一致,避免空字符串拼出裸后缀。
    for prefix_field, prefix_value in (
        ("ssid_24_prefix", template.ssid_24_prefix),
        ("ssid_5_prefix", template.ssid_5_prefix),
    ):
        if not (prefix_value or "").strip():
            raise ValueError("2.4G SSID 和 5G SSID 前缀均不能为空。")
    records: list[SampleParameterRecord] = []
    for index, vlan in enumerate(normalized_vlans, start=2):
        values = {column: str(template.fixed_values.get(column, "")).strip() for column in PARAMETER_COLUMNS}
        values["SN"] = _unique_numeric_sn(template.sn_prefix, suffix_length, seen_values["SN"])
        values["GPON SN"] = _unique_fixed_xpon_sn(template.xpon_prefix, seen_values["GPON SN"])
        # 2026-07-19 优化:2.4G / 5G SSID 后缀改为本机 MAC 末 4 位十六进制
        # (与 MAC 强绑定,同一行两个 SSID 自动共享后缀,操作员可一眼对应本机 MAC)。
        # 例:MAC=CA:7A:85:18:B0:E8,后缀="B0E8" → SSID = prefix + "B0E8"。
        mac = _unique_mac(mac_prefix, seen_mac)
        ssid_suffix = mac.replace(":", "")[-4:]
        values["MAC"] = mac
        values["2.4G SSID"] = f"{template.ssid_24_prefix}{ssid_suffix}"
        values["5G SSID"] = f"{template.ssid_5_prefix}{ssid_suffix}"
        records.append(SampleParameterRecord(vlan=vlan, values=values, source_row=index))
    return records


def _unique_fixed_xpon_sn(prefix: str, seen: set[str]) -> str:
    """生成固定12位XPON SN：4位操作员前缀 + 8位大写十六进制后缀。"""
    clean_prefix = (prefix or "").strip()
    if len(clean_prefix) != XPON_SN_PREFIX_LENGTH:
        raise ValueError(
            f"XPON SN 前缀必须固定为 {XPON_SN_PREFIX_LENGTH} 位，生成结果才是 "
            f"{XPON_SN_TOTAL_LENGTH} 位（后 {XPON_SN_HEX_SUFFIX_LENGTH} 位为十六进制）。"
        )
    for _ in range(1000):
        suffix = "".join(
            secrets.choice(_UPPER_HEX_ALPHABET)
            for _ in range(XPON_SN_HEX_SUFFIX_LENGTH)
        )
        value = clean_prefix + suffix
        if value not in seen:
            seen.add(value)
            return value
    raise RuntimeError("随机 XPON SN 生成碰撞次数过多，请重新生成。")


def _unique_numeric_sn(prefix: str, suffix_length: int, seen: set[str]) -> str:
    """生成纯数字SN，前缀与随机后缀合计最多16位。"""
    clean_prefix = (prefix or "").strip()
    if not clean_prefix:
        raise ValueError("SN 前缀不能为空。")
    if not clean_prefix.isascii() or not clean_prefix.isdigit():
        raise ValueError("SN 前缀必须为纯数字（仅 0-9）。")
    total_length = len(clean_prefix) + suffix_length
    if total_length > SN_MAX_LENGTH:
        raise ValueError(
            f"SN 前缀长度（{len(clean_prefix)}）+ 随机后缀长度（{suffix_length}）"
            f"不能超过 {SN_MAX_LENGTH} 位。"
        )
    for _ in range(1000):
        value = clean_prefix + "".join(
            secrets.choice(_DIGIT_ALPHABET) for _ in range(suffix_length)
        )
        if value not in seen:
            seen.add(value)
            return value
    raise RuntimeError("随机 SN 生成碰撞次数过多，请重新生成。")


def _normalize_mac_prefix(prefix: str) -> bytes:
    """校验用户输入的MAC前6位，并返回3字节前缀。"""
    compact = (prefix or "").strip().replace(":", "").replace("-", "")
    if len(compact) != 6:
        raise ValueError("MAC 前6位必须包含恰好 6 个十六进制字符。")
    if not compact.isascii() or any(character not in _UPPER_HEX_ALPHABET for character in compact.upper()):
        raise ValueError("MAC 前6位必须为十六进制字符（仅 0-9、A-F）。")
    raw = bytes.fromhex(compact)
    if raw[0] & 0x01:
        raise ValueError("MAC 前6位的首字节必须为单播地址（最低位不能为 1）。")
    return raw


def _unique_mac(prefix: bytes, seen: set[str]) -> str:
    """根据用户指定的3字节前缀生成唯一的3字节随机后缀。"""
    for _ in range(1000):
        nic = bytearray(secrets.token_bytes(3))
        raw = prefix + bytes(nic)
        value = ":".join(f"{item:02X}" for item in raw)
        if value not in seen:
            seen.add(value)
            return value
    raise RuntimeError("随机 MAC 生成碰撞次数过多，请重新生成。")


def _read_rows(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("读取 Excel 参数表需要安装 openpyxl。") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("读取 .xls 参数表需要安装 xlrd。") from exc
        workbook = xlrd.open_workbook(path)
        sheet = workbook.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return [list(row) for row in csv.reader(stream)]
    except UnicodeDecodeError:
        with path.open("r", encoding="gbk", newline="") as stream:
            return [list(row) for row in csv.reader(stream)]


def _cell_for(header: list[str], row: list[str], column: str, aliases: tuple[str, ...] = ()) -> str:
    for candidate in (column, *aliases):
        indices = _indices_for(header, candidate)
        if indices:
            index = indices[0]
            return row[index].strip() if index < len(row) else ""
    return ""


def _indices_for(header: list[str], target: str) -> list[int]:
    normalized = target.strip().casefold()
    return [index for index, value in enumerate(header) if value.strip().casefold() == normalized]
